"""Real parser behavior on synthetic files, including unsupported content."""

import io
import zipfile

import pytest

from telegram_mcp.adapters.extraction import parsers
from telegram_mcp.domain.models import ToolError
from tests.support.extraction import docx_bytes, pdf_bytes


def test_pdf_text_and_physical_pages_survive_blank_pages() -> None:
    result = parsers.extract_document(pdf_bytes(["Café", None, "Last page"]))
    assert result.page_count == 3 and result.text == "Café\n\n\n\nLast page"
    assert [span.page_number for span in result.page_spans] == [1, 2, 3]
    assert result.issues[0].code == "DOCUMENT_PAGES_WITHOUT_TEXT"
    assert "preview: 2" in result.issues[0].message


@pytest.mark.parametrize("pages", [[], [None], [None, None]])
def test_no_text_does_not_claim_the_file_is_a_scan(pages: list) -> None:
    with pytest.raises(ToolError, match="DOCUMENT_NO_TEXT"):
        parsers.extract_document(pdf_bytes(pages))


def test_encrypted_pdf_does_not_accept_a_password() -> None:
    with pytest.raises(ToolError, match="DOCUMENT_ENCRYPTED"):
        parsers.extract_document(pdf_bytes(["secret"], encrypted=True))


def test_page_and_text_budgets_publish_no_partial_document(monkeypatch) -> None:
    monkeypatch.setattr(parsers, "MAX_PDF_PAGES", 1)
    with pytest.raises(ToolError, match="EXTRACTION_LIMIT_EXCEEDED"):
        parsers.extract_document(pdf_bytes(["one", "two"]))
    monkeypatch.setattr(parsers, "MAX_TEXT_CHARACTERS", 3)
    with pytest.raises(ToolError, match="EXTRACTION_LIMIT_EXCEEDED"):
        parsers.extract_document(docx_bytes("four"))


def test_docx_unicode_is_read_without_page_numbers() -> None:
    result = parsers.extract_document(docx_bytes())
    assert result.text == "Résumé 😀\nПривет 世界\n"
    assert result.format == "docx" and result.page_count is None and result.page_spans == []


def test_docx_table_cells_use_tabs_and_rows_use_newlines() -> None:
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip((cell for row in table.rows for cell in row.cells), "ABCD", strict=True):
        cell.text = text
    output = io.BytesIO()
    doc.save(output)
    assert parsers.extract_document(output.getvalue()).text == "A\tB\nC\tD\n"


def test_docx_preserves_tables_links_and_omission_notice() -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph("Before ")
    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"),
        doc.part.relate_to("https://example.invalid/no-fetch", RT.HYPERLINK, is_external=True),
    )
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "visible link"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Horizontal merge"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Vertical merge"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    doc.add_paragraph("After")
    doc.sections[0].header.paragraphs[0].text = "Omitted header"
    output = io.BytesIO()
    doc.save(output)
    result = parsers.extract_document(output.getvalue())
    assert (
        result.text.index("Before visible link")
        < result.text.index("Nested")
        < result.text.index("After")
    )
    assert result.text.count("Horizontal merge") == result.text.count("Vertical merge") == 1
    assert "Omitted header" not in result.text
    assert [issue.code for issue in result.issues] == ["DOCUMENT_CONTENT_OMITTED"]


def test_docx_expansion_budget_is_checked_before_loading(monkeypatch) -> None:
    data = docx_bytes("a" * 1000)
    monkeypatch.setattr(parsers, "MAX_DOCX_BYTES", 100)
    with pytest.raises(ToolError, match="EXTRACTION_LIMIT_EXCEEDED"):
        parsers.extract_document(data)


def test_damaged_docx_xml_returns_a_sanitized_document_error() -> None:
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(docx_bytes())) as original,
        zipfile.ZipFile(output, "w") as damaged,
    ):
        for name in original.namelist():
            data = (
                b"<unclosed>private-canary" if name == "word/document.xml" else original.read(name)
            )
            damaged.writestr(name, data)
    with pytest.raises(ToolError, match="DOCUMENT_INVALID") as caught:
        parsers.extract_document(output.getvalue())
    assert "private-canary" not in caught.value.detail


def test_an_arbitrary_zip_is_not_a_docx() -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("notes.txt", "text")
    with pytest.raises(ToolError, match="DOCUMENT_UNSUPPORTED"):
        parsers.extract_document(output.getvalue())


@pytest.mark.parametrize("data", [b"plain text", bytes.fromhex("d0cf11e0a1b11ae1") + b"Office"])
def test_unsupported_content_is_distinct_from_damage(data: bytes) -> None:
    with pytest.raises(ToolError, match="DOCUMENT_UNSUPPORTED"):
        parsers.extract_document(data)
