"""Synthetic documents and isolated extraction contexts without a Telegram account."""

import asyncio
import io
import re
import zipfile
from collections.abc import AsyncGenerator
from contextlib import asynccontextmanager
from dataclasses import replace
from pathlib import Path

from telegram_mcp.adapters.clock import SystemClock
from telegram_mcp.adapters.config import Settings
from telegram_mcp.adapters.downloads import DownloadStore
from telegram_mcp.adapters.extraction.store import ExtractionStore
from telegram_mcp.application.service import Service
from telegram_mcp.bootstrap import build_service
from telegram_mcp.domain.extraction import ExtractedDocument
from telegram_mcp.domain.models import Chat, ToolError
from tests.support.downloads import AttachmentBackend
from tests.support.fakes import FakeClock


def pdf_bytes(pages: list[str | None], *, encrypted: bool = False) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=200, height=200)
        if text is not None:
            font = DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                    NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
                }
            )
            page[NameObject("/Resources")] = DictionaryObject(
                {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
            )
            stream = DecodedStreamObject()
            encoded = text.encode("cp1252").hex().encode()
            stream.set_data(b"BT /F1 12 Tf 10 100 Td <" + encoded + b"> Tj ET")
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("test-password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(text: str = "Résumé 😀\nПривет 世界") -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def damaged_pdf_bytes(defect: str) -> bytes:
    """Small recoverable defects, including a stream failure after a valid first page."""
    data = pdf_bytes(["First", None, "Last"])
    if defect == "xref":
        return re.sub(rb"(startxref\s+)(\d+)", lambda m: m[1] + str(int(m[2]) + 1).encode(), data)
    if defect == "offsets":
        return re.sub(
            rb"(?m)^(\d{10})( \d{5} n )",
            lambda m: f"{int(m[1]) + 1:010}".encode() + m[2],
            data,
        )
    if defect == "length":
        match = list(re.finditer(rb"/Length (\d+)", data))[-1]
        return data[: match.start(1)] + b"1".rjust(len(match[1]), b"0") + data[match.end(1) :]
    if defect == "prefix":
        return b"\xef\xbb\xbf" + data
    raise ValueError(f"Unknown PDF defect: {defect}")


def relocated_docx_bytes(main_name: str = "word/document2.xml", text: str = "Relocated") -> bytes:
    """Move the main part and its relationship file, preserving OPC references."""
    output = io.BytesIO()
    main_relationships = "word/_rels/" + main_name.removeprefix("word/") + ".rels"
    with (
        zipfile.ZipFile(io.BytesIO(docx_bytes(text))) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            data = source.read(name)
            if name in ("[Content_Types].xml", "_rels/.rels"):
                data = data.replace(b"word/document.xml", main_name.encode())
            name = {
                "word/document.xml": main_name,
                "word/_rels/document.xml.rels": main_relationships,
            }.get(name, name)
            target.writestr(name, data)
    return output.getvalue()


class FakeParser:
    def __init__(self) -> None:
        self.available = True
        self.document = ExtractedDocument(format="docx", text="Résumé 😀\nПривет 世界\n" * 4)
        self.calls = 0
        self.error: ToolError | None = None
        self.entered = asyncio.Event()
        self.resume: asyncio.Event | None = None
        self.closed = False

    async def extract(self, data: bytes) -> ExtractedDocument:
        self.calls += 1
        self.entered.set()
        if self.resume is not None:
            await self.resume.wait()
        if self.error is not None:
            raise self.error
        return self.document

    async def close(self) -> None:
        self.closed = True


def extraction_backend(*, recovered_pdf: bool = False) -> AttachmentBackend:
    backend = AttachmentBackend()
    backend.chats = [Chat(id="1", title="Synthetic documents", type="group")]
    backend.data = damaged_pdf_bytes("length") if recovered_pdf else docx_bytes()
    backend.source = replace(
        backend.source,
        size=len(backend.data),
        name="selected.pdf" if recovered_pdf else "selected.docx",
        mime_type=(
            "application/pdf"
            if recovered_pdf
            else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
    return backend


def extraction_stdio_service(path: Path, *, recovered_pdf: bool = False) -> Service:
    settings = Settings(
        allowed_chat_ids=["1"], allow_attachment_downloads=True, session_path=path / "session"
    )
    clock = SystemClock()
    return build_service(
        extraction_backend(recovered_pdf=recovered_pdf),
        settings,
        clock=clock,
        downloads=DownloadStore(settings.session_path, clock, root=path / "downloads"),
        extractions=ExtractionStore(settings.session_path, clock, root=path / "extractions"),
    )


@asynccontextmanager
async def extraction_context(
    tmp_path: Path,
) -> AsyncGenerator[
    tuple[Service, AttachmentBackend, Settings, FakeParser, ExtractionStore, FakeClock]
]:
    clock, backend, parser = FakeClock(), AttachmentBackend(), FakeParser()
    backend.source = replace(backend.source, name="selected.docx")
    settings = Settings(
        allowed_chat_ids=["1"],
        allow_attachment_downloads=True,
        session_path=tmp_path / "account.session",
    )
    store = ExtractionStore(settings.session_path, clock, root=tmp_path / "extractions")
    service = build_service(
        backend,
        settings,
        clock=clock,
        document_parser=parser,
        extractions=store,
        downloads=DownloadStore(settings.session_path, clock, root=tmp_path / "downloads"),
    )
    try:
        yield service, backend, settings, parser, store, clock
    finally:
        await service.close()
