"""Bounded parsing of selected bytes; imported only inside an extraction worker."""

import gc
import io
import zipfile
from collections.abc import Iterator
from pathlib import PurePosixPath
from typing import Any

from telegram_mcp.application.extraction_ports import (
    MAX_DOCX_BYTES,
    MAX_DOCX_ENTRIES,
    MAX_DOCX_ENTRY_BYTES,
)
from telegram_mcp.domain.extraction import (
    MAX_PDF_PAGES,
    MAX_TEXT_CHARACTERS,
    ExtractedDocument,
    PageSpan,
)
from telegram_mcp.domain.models import Issue, ToolError


def limit_error(detail: str) -> ToolError:
    return ToolError("EXTRACTION_LIMIT_EXCEEDED", detail)


class TextBuilder:
    def __init__(self) -> None:
        self.parts: list[str] = []
        self.length = 0

    def append(self, text: str) -> None:
        if self.length + len(text) > MAX_TEXT_CHARACTERS:
            raise limit_error("Extracted text exceeds 1,000,000 Unicode characters.")
        self.parts.append(text)
        self.length += len(text)

    def finish(self) -> str:
        text = "".join(self.parts)
        if not text.strip():
            raise ToolError(
                "DOCUMENT_NO_TEXT",
                "No extractable text was found in the supported document content. "
                "An image-only document may require OCR, which is not included.",
            )
        return text


def extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf.errors import FileNotDecryptedError, PdfReadError

    try:
        return _read_pdf(data, strict=True)
    except FileNotDecryptedError:
        raise
    except PdfReadError:
        pass
    # Release cyclic reader/page references from the failed attempt before retrying.
    gc.collect()
    document = _read_pdf(data, strict=False)
    document.issues.append(
        Issue(
            code="DOCUMENT_RECOVERED",
            message="The PDF required tolerant parsing. Text recovery is best effort; "
            "complete recovery is not guaranteed.",
        )
    )
    return document


def _read_pdf(data: bytes, *, strict: bool) -> ExtractedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=strict)
    if reader.is_encrypted:
        raise ToolError(
            "DOCUMENT_ENCRYPTED", "Encrypted PDFs are not supported; no password is accepted."
        )
    count = len(reader.pages)
    if count > MAX_PDF_PAGES:
        raise limit_error("PDF documents are limited to 500 pages.")
    builder, spans, empty = TextBuilder(), [], []
    for index, page in enumerate(reader.pages):
        start = builder.length
        text = page.extract_text() or ""
        if not text.strip():
            empty.append(index + 1)
        builder.append(text)
        if index + 1 < count:
            builder.append("\n\n")
        spans.append(PageSpan(page_number=index + 1, start=start, end=builder.length))
    text = builder.finish()
    issues = []
    if empty:
        preview = ", ".join(map(str, empty[:20]))
        issues.append(
            Issue(
                code="DOCUMENT_PAGES_WITHOUT_TEXT",
                message=f"{len(empty)} page(s) have no extractable text; page preview: {preview}. "
                "These may be blank or contain images; no OCR was performed.",
            )
        )
    return ExtractedDocument(
        format="pdf", text=text, page_count=count, page_spans=spans, issues=issues
    )


def check_package(data: bytes) -> set[str]:
    """Validate ZIP sizes using actual bounded decompression, without extracting paths."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES or sum(e.file_size for e in entries) > MAX_DOCX_BYTES:
            raise limit_error("DOCX packages are limited to 64 MiB expanded and 2048 entries.")
        names = {entry.filename for entry in entries}
        if not {"[Content_Types].xml", "_rels/.rels"} <= names:
            raise ToolError("DOCUMENT_UNSUPPORTED", "The ZIP file is not a supported DOCX package.")
        if len(names) != len(entries):
            raise ToolError("DOCUMENT_INVALID", "The DOCX package has duplicate entries.")
        for entry in entries:
            _check_entry(entry)
        expanded = 0
        for entry in entries:
            expanded += _expanded_size(archive, entry, MAX_DOCX_BYTES - expanded)
        return names


def _check_entry(entry: zipfile.ZipInfo) -> None:
    path = PurePosixPath(entry.filename)
    if path.is_absolute() or ".." in path.parts or "\\" in entry.filename:
        raise ToolError("DOCUMENT_INVALID", "The DOCX package has invalid entry paths.")
    if entry.flag_bits & 1:
        raise ToolError("DOCUMENT_ENCRYPTED", "Encrypted document packages are unsupported.")
    if entry.file_size > MAX_DOCX_ENTRY_BYTES:
        raise limit_error("Each expanded DOCX entry is limited to 32 MiB.")


def _expanded_size(archive: zipfile.ZipFile, entry: zipfile.ZipInfo, remaining: int) -> int:
    expanded = 0
    with archive.open(entry) as stream:
        while chunk := stream.read(
            min(65536, remaining - expanded + 1, MAX_DOCX_ENTRY_BYTES - expanded + 1)
        ):
            expanded += len(chunk)
            if expanded > MAX_DOCX_ENTRY_BYTES:
                raise limit_error("Each expanded DOCX entry is limited to 32 MiB.")
            if expanded > remaining:
                raise limit_error("Expanded DOCX content exceeds 64 MiB.")
    return expanded


def _blocks(container: Any, depth: int = 0) -> Iterator[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    if depth > 64:
        raise limit_error("DOCX table nesting exceeds 64 levels.")
    for index, block in enumerate(container.iter_inner_content()):
        if index:
            yield "\n"
        if isinstance(block, Paragraph):
            yield block.text
        elif isinstance(block, Table):
            yield from _table(block, depth)


def _table(table: Any, depth: int) -> Iterator[str]:
    # python-docx repeats merged cells in its rectangular grid. Keep their XML
    # identities to include the underlying cell once, including vertical merges.
    seen = set()
    for row_index, row in enumerate(table.rows):
        if row_index:
            yield "\n"
        for index, cell in enumerate(row.cells):
            if index:
                yield "\t"
            if cell._tc not in seen:
                seen.add(cell._tc)
                yield from _blocks(cell, depth + 1)


def _omissions(document: Any, names: set[str]) -> list[Issue]:
    omitted = set()
    for name in names:
        if name.startswith(("word/header", "word/footer")):
            omitted.add("headers/footers")
        if name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
            omitted.add("notes/comments")
    unsupported = {
        "txbxContent",
        "ins",
        "del",
        "moveFrom",
        "moveTo",
        "sdt",
        "altChunk",
        "oMath",
        "fldSimple",
        "customXml",
        "AlternateContent",
    }
    if any(
        isinstance(node.tag, str) and node.tag.rsplit("}", 1)[-1] in unsupported
        for node in document.element.iter()
    ):
        omitted.add("text boxes, revisions, fields or unsupported body elements")
    if not omitted:
        return []
    return [
        Issue(
            code="DOCUMENT_CONTENT_OMITTED",
            message="Extracted main-body paragraphs and tables only; omitted: "
            + ", ".join(sorted(omitted))
            + ".",
        )
    ]


def extract_docx(data: bytes) -> ExtractedDocument:
    from docx import Document

    names = check_package(data)
    try:
        document = Document(io.BytesIO(data))
    except ValueError as exc:
        if "is not a Word file, content type is" in str(exc):
            raise ToolError(
                "DOCUMENT_UNSUPPORTED", "The package's main part is not a supported DOCX document."
            ) from None
        raise
    builder = TextBuilder()
    for text in _blocks(document):
        builder.append(text)
    builder.append("\n")
    return ExtractedDocument(
        format="docx", text=builder.finish(), issues=_omissions(document, names)
    )


def extract_document(data: bytes) -> ExtractedDocument:
    from docx.opc.exceptions import OpcError

    # lxml's binary extension does not ship type stubs.
    from lxml import etree  # ty: ignore[unresolved-import]
    from pypdf.errors import FileNotDecryptedError, LimitReachedError, PyPdfError

    try:
        return _detect_and_extract(data)
    except FileNotDecryptedError:
        raise ToolError("DOCUMENT_ENCRYPTED", "Encrypted PDFs are not supported.") from None
    except LimitReachedError:
        raise limit_error("The PDF parser reached its processing limit.") from None
    except etree.XMLSyntaxError as exc:
        if any(entry.type == etree.ErrorTypes.ERR_NO_MEMORY for entry in exc.error_log):
            raise limit_error("The XML parser could not allocate enough memory.") from None
        raise ToolError(
            "DOCUMENT_INVALID", "The document is damaged or cannot be parsed."
        ) from None
    except (
        PyPdfError,
        OpcError,
        SyntaxError,
        zipfile.BadZipFile,
        ValueError,
        KeyError,
        IndexError,
        EOFError,
        RecursionError,
    ):
        raise ToolError(
            "DOCUMENT_INVALID", "The document is damaged or cannot be parsed."
        ) from None


def _detect_and_extract(data: bytes) -> ExtractedDocument:
    if data.startswith(b"PK"):
        return extract_docx(data)
    if b"%PDF-" in data[:1024]:
        return extract_pdf(data)
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise ToolError(
            "DOCUMENT_UNSUPPORTED",
            "Select an unencrypted DOCX package. Encrypted Office "
            "containers and legacy .doc files are not supported.",
        )
    raise ToolError(
        "DOCUMENT_UNSUPPORTED", "Only text-based PDF and DOCX attachments are supported."
    )
